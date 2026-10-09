import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from fpdf import FPDF
from jinja2 import Template
from typing import Dict, Any

class ReportEngine:
    def __init__(self, data: Dict[str, Any], insights: str = None):
        self.data = data
        self.insights = insights

    def generate_pptx(self, output_path: str):
        """Generates a professional PowerPoint presentation."""
        prs = Presentation()
        
        # Slide 1: Title
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Analytics Pro Dashboard"
        subtitle.text = "Relatório Executivo de Performance"

        # Slide 2: Executive Summary
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Resumo Executivo"
        
        tf = slide.placeholders[1].text_frame
        summary = self.data.get('summary', {})
        tf.text = f"Valor Total: R$ {summary.get('total_value', 0):,.2f}"
        p = tf.add_paragraph()
        p.text = f"Ticket Médio: R$ {summary.get('avg_transaction', 0):,.2f}"
        p = tf.add_paragraph()
        p.text = f"Total de Fornecedores: {summary.get('count_suppliers', 0)}"

        # Slide 3: Top Suppliers
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Top 10 Fornecedores"
        
        suppliers = self.data.get('suppliers', {})
        rows = len(suppliers) + 1
        cols = 2
        left = Inches(1)
        top = Inches(2)
        width = Inches(8)
        height = Inches(0.5)
        
        table = slide.shapes.add_table(rows, cols, left, top, width, height).table
        table.columns[0].width = Inches(5)
        table.columns[1].width = Inches(3)
        
        table.cell(0, 0).text = "Fornecedor"
        table.cell(0, 1).text = "Valor Total"
        
        for i, (name, val) in enumerate(suppliers.items(), start=1):
            table.cell(i, 0).text = str(name)
            table.cell(i, 1).text = f"R$ {val:,.2f}"

        # Slide 4: Insights Estratégicos (IA)
        if self.insights:
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = "Insights Estratégicos (IA)"
            tf = slide.placeholders[1].text_frame
            tf.word_wrap = True
            
            lines = self.insights.split('\n')
            first = True
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                # Extract bullet content if line has bullet format
                if line.startswith(('-', '*', '1.', '2.', '3.', '4.', '5.')):
                    clean_line = re.sub(r'^[\-\*\d\.\s]+', '', line).strip()
                    clean_line = clean_line.replace('**', '')
                    if first:
                        tf.text = clean_line
                        first = False
                    else:
                        p = tf.add_paragraph()
                        p.text = clean_line
                        p.level = 0
                elif line.startswith('###') or line.startswith('##'):
                    clean_line = line.replace('###', '').replace('##', '').strip().replace('**', '')
                    if first:
                        tf.text = clean_line
                        first = False
                    else:
                        p = tf.add_paragraph()
                        p.text = clean_line
                        p.font.bold = True
                        p.level = 0

        prs.save(output_path)
        return output_path

    def generate_pdf(self, output_path: str):
        """Generates a structured PDF report."""
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, "Relatório Analítico de Operações", ln=True, align='C')
        pdf.ln(10)
        
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, "Resumo Geral", ln=True)
        pdf.set_font("Arial", '', 10)
        summary = self.data.get('summary', {})
        pdf.cell(0, 8, f"Valor Total Transacionado: R$ {summary.get('total_value', 0):,.2f}", ln=True)
        pdf.cell(0, 8, f"Quantidade de Transações: {summary.get('count_transactions', 0)}", ln=True)
        pdf.ln(5)
        
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, "Distribuição por Operação", ln=True)
        pdf.set_font("Arial", '', 10)
        ops = self.data.get('operations', {}).get('by_operation', {})
        for op, val in ops.items():
            pdf.cell(0, 8, f"{op}: R$ {val:,.2f}", ln=True)
            
        if self.insights:
            pdf.ln(10)
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(0, 10, "Insights Estratégicos (IA)", ln=True)
            pdf.set_font("Arial", '', 9)
            # Limpa cabeçalhos markdown e negritos para texto corrido no PDF
            clean_text = self.insights.replace("###", "").replace("##", "").replace("**", "")
            pdf.multi_cell(0, 6, clean_text)
            
        pdf.output(output_path)
        return output_path

    def generate_html(self, output_path: str):
        """Generates a standalone HTML report."""
        template_str = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>Relatório Web - Analytics Pro</title>
            <style>
                body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; margin: 40px; color: #333; }
                h1 { color: #4F46E5; }
                .card { border: 1px solid #ddd; padding: 20px; border-radius: 8px; margin-bottom: 20px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }
                .grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 20px; }
                table { width: 100%; border-collapse: collapse; margin-top: 20px; }
                th, td { border: 1px solid #ddd; padding: 12px; text-align: left; }
                th { background-color: #f3f4f6; }
            </style>
        </head>
        <body>
            <h1>Analytics Pro Dashboard - Relatório Interativo</h1>
            <div class="grid">
                <div class="card">
                    <h3>Valor Total</h3>
                    <p>R$ {{ data.summary.total_value | round(2) }}</p>
                </div>
                <div class="card">
                    <h3>Total Fornecedores</h3>
                    <p>{{ data.summary.count_suppliers }}</p>
                </div>
            </div>
            
            {% if insights %}
            <div class="card" style="border-left: 4px solid #4F46E5; background-color: #f5f3ff;">
                <h3 style="color: #4F46E5; margin-top: 0;">Insights Estratégicos (IA)</h3>
                <div style="white-space: pre-wrap; font-size: 0.95rem; line-height: 1.6;">{{ insights }}</div>
            </div>
            {% endif %}
            
            <div class="card">
                <h3>Top 10 Fornecedores</h3>
                <table>
                    <thead><tr><th>Fornecedor</th><th>Valor</th></tr></thead>
                    <tbody>
                        {% for name, val in data.suppliers.items() %}
                        <tr><td>{{ name }}</td><td>R$ {{ val | round(2) }}</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </body>
        </html>
        """
        template = Template(template_str)
        html_content = template.render(data=self.data, insights=self.insights)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html_content)
        return output_path
